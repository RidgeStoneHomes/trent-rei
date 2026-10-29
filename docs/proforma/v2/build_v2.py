import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
from uses import U, COLS, KEYS

NAVY = "1F3864"; YEL = "FFF2CC"; F = "Arial"
fn = Font(name=F, size=10); fin = Font(name=F, size=10, color="0000FF"); flink = Font(name=F, size=10, color="008000")
fb = Font(name=F, size=10, bold=True); fh = Font(name=F, size=10, bold=True, color="FFFFFF")
ft = Font(name=F, size=14, bold=True, color=NAVY); fs = Font(name=F, size=10, italic=True, color="595959")
fsec = Font(name=F, size=11, bold=True, color=NAVY)
fill_h = PatternFill("solid", fgColor=NAVY); fill_in = PatternFill("solid", fgColor=YEL)
fill_t = PatternFill("solid", fgColor="F2F2F2"); fill_g = PatternFill("solid", fgColor="FBF3D5")
USD = '$#,##0;($#,##0);"-"'; USD2 = '$#,##0.00;($#,##0.00);"-"'; NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'; NUM3 = '#,##0.000'; PCT = '0.0%;(0.0%);"-"'; MULT = '0.00"x"'

wb = Workbook()

def title(ws, t, s):
    ws["A1"] = t; ws["A1"].font = ft; ws["A2"] = s; ws["A2"].font = fs

def hdr(ws, row, labels, start=1, height=None):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=l); c.font = fh; c.fill = fill_h
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if height: ws.row_dimensions[row].height = height

def put(ws, ref, v, font=fn, fmt=None, fill=None):
    c = ws[ref]; c.value = v; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    return c

def dv(ws, ref, formula):
    d = DataValidation(type="list", formula1=formula, allow_blank=True); ws.add_data_validation(d); d.add(ref)

# ============================================================ Lookups
lk = wb.active; lk.title = "Lookups"
title(lk, "Lookups", "Dropdown lists and adjustment factors. Factors are placeholders - calibrate from the cost library.")
hdr(lk, 4, ["County", "Regional cost factor"], 1)
for i, (c, f) in enumerate([("Grayson", 1.0), ("Collin", 1.12), ("Denton", 1.10), ("Cooke", 0.97), ("Fannin", 0.96), ("Other TX", 1.0)]):
    put(lk, f"A{5+i}", c); put(lk, f"B{5+i}", f, fin, '0.00')
hdr(lk, 4, ["Septic soil", "Septic multiplier"], 4)
for i, (s, m) in enumerate([("Good", 1.0), ("Fair", 1.2), ("Poor", 1.5)]):
    put(lk, f"D{5+i}", s); put(lk, f"E{5+i}", m, fin, '0.00')
hdr(lk, 4, ["Max slope", "Grading multiplier"], 7)
for i, (s, m) in enumerate([(0.05, 1.0), (0.10, 1.25), (1.0, 1.6)]):
    put(lk, f"G{5+i}", s, fin, PCT); put(lk, f"H{5+i}", m, fin, '0.00')
hdr(lk, 4, ["Scenario"], 10)
for i, s in enumerate(["Low", "Base", "High"]): put(lk, f"J{5+i}", s)
hdr(lk, 4, ["Y/N"], 12); put(lk, "L5", "Y"); put(lk, "L6", "N")
hdr(lk, 4, ["Engine"], 14)
for i, s in enumerate(["For-Sale", "Income", "Operator Pad", "Land Play"]): put(lk, f"N{5+i}", s)
for k, w in {"A": 12, "B": 20, "D": 14, "E": 18, "G": 11, "H": 18, "J": 11, "L": 6, "N": 14}.items():
    lk.column_dimensions[k].width = w

# ============================================================ Parcel
pc = wb.create_sheet("Parcel", 0)
title(pc, "Parcel Inputs", "One parcel. In the software these fill from Landova's site analysis. Blue on yellow = input.")
hdr(pc, 4, ["Input", "Value", "Unit", "Source in Landova"], 2)
P = {}
prow = [
    ("name", "Parcel name", "Sample Parcel - Grayson County", "", "User", None),
    ("county", "County", "Grayson", "", "Parcel record", None),
    ("scenario", "Cost scenario", "Base", "Low/Base/High", "User", None),
    ("gross", "Gross acres", 20, "ac", "Parcel boundary", NUM1),
    ("buildpct", "Buildable %", 0.78, "% of gross", "Buildable-land analysis", PCT),
    ("slope", "Average slope", 0.04, "%", "Topography", PCT),
    ("tree", "Tree cover", 0.35, "%", "3D tree coverage", PCT),
    ("frontage", "Road frontage", 660, "LF", "Frontage analysis", NUM),
    ("water_site", "Public water at site", "Y", "Y/N", "Utilities layer", None),
    ("water_dist", "Distance to water main", 0, "LF", "Utilities layer", NUM),
    ("sewer_site", "Public sewer at site", "N", "Y/N", "Utilities layer", None),
    ("sewer_dist", "Distance to sewer main", 6000, "LF", "Utilities layer", NUM),
    ("soil", "Septic soil rating", "Fair", "Good/Fair/Poor", "Soil composition", None),
    ("land", "Asking / offer price (whole parcel)", 300000, "$", "Listing or user", USD),
    ("vpd", "Traffic count on frontage road", 4000, "vehicles/day", "TxDOT counts", NUM),
    ("hh", "Households within 3 miles", 2500, "households", "Census / growth layer", NUM),
    ("hwy", "Miles to nearest highway interchange", 6, "miles", "Road network", NUM1),
    ("corner", "Corner lot (hard corner on two roads)", "N", "Y/N", "Frontage analysis", None),
    ("sub", "Miles to substation with capacity", 3, "miles", "Utility / transmission layer", NUM1),
]
for i, (k, lab, v, unit, src, fmt) in enumerate(prow):
    r = 5 + i; P[k] = f"Parcel!$C${r}"
    put(pc, f"B{r}", lab); put(pc, f"C{r}", v, fin, fmt, fill_in); put(pc, f"D{r}", unit); put(pc, f"E{r}", src, fs)
dv(pc, "C6", "=Lookups!$A$5:$A$10"); dv(pc, "C7", "=Lookups!$J$5:$J$7"); dv(pc, "C13", "=Lookups!$L$5:$L$6")
dv(pc, "C15", "=Lookups!$L$5:$L$6"); dv(pc, "C17", "=Lookups!$D$5:$D$7"); dv(pc, "C22", "=Lookups!$L$5:$L$6")
r = 5 + len(prow) + 1
put(pc, f"B{r}", "Derived (formulas)", fsec); r += 1
derived = [
    ("net", "Net buildable acres", f"={P['gross']}*{P['buildpct']}", NUM1),
    ("pubW", "Public water used (1 = yes)", None, NUM),
    ("pubS", "Public sewer used (1 = yes)", None, NUM),
    ("extW", "Off-site water extension LF", None, NUM),
    ("extS", "Off-site sewer extension LF", None, NUM),
    ("sidx", "Scenario column", f"=MATCH({P['scenario']},Lookups!$J$5:$J$7,0)", NUM),
    ("reg", "Regional cost factor", f"=IFERROR(INDEX(Lookups!$B$5:$B$10,MATCH({P['county']},Lookups!$A$5:$A$10,0)),1)", '0.00'),
    ("slopem", "Grading multiplier (slope)", f"=INDEX(Lookups!$H$5:$H$7,MATCH(TRUE,INDEX({P['slope']}<=Lookups!$G$5:$G$7,0),0))", '0.00'),
    ("soilm", "Septic multiplier (soil)", f"=IFERROR(INDEX(Lookups!$E$5:$E$7,MATCH({P['soil']},Lookups!$D$5:$D$7,0)),1)", '0.00'),
]
for k, lab, f, fmt in derived:
    P[k] = f"Parcel!$C${r}"; put(pc, f"B{r}", lab); put(pc, f"C{r}", f, fn, fmt); r += 1
pc.column_dimensions["A"].width = 2; pc.column_dimensions["B"].width = 40; pc.column_dimensions["C"].width = 30
pc.column_dimensions["D"].width = 14; pc.column_dimensions["E"].width = 34

# ============================================================ Assumptions
asn = wb.create_sheet("Assumptions", 1)
title(asn, "Global Assumptions", "Apply to every use. Placeholders - set to your lender terms and market.")
hdr(asn, 4, ["Assumption", "Value", "Notes"], 2)
A = {}
alist = [
    ("wmax", "Max water main extension before wells", 2640, NUM, "LF"),
    ("smax", "Max sewer main extension before septic", 1320, NUM, "LF"),
    ("det", "Detention % of net", 0.04, PCT, "Confirm with drainage study"),
    ("rowpct", "ROW share of lot area (septic lot rule)", 0.15, PCT, ""),
    ("lot_w", "Min lot with septic + public water", 0.5, NUM1, "acres; TX rule of thumb - confirm by county"),
    ("lot_well", "Min lot with septic + well", 1.0, NUM1, "acres; confirm by county"),
    ("hcont", "Horizontal contingency", 0.10, PCT, "% of horizontal costs"),
    ("vsoft", "Vertical soft costs (A&E, permits, dev fee)", 0.08, PCT, "% of vertical"),
    ("vcont", "Vertical contingency", 0.05, PCT, "% of vertical"),
    ("closing", "Closing costs on land", 0.02, PCT, "% of land price"),
    ("carry", "Land carry (tax + insurance) per year", 0.025, PCT, "% of land price"),
    ("sell_fs", "Selling costs - homes / condos", 0.06, PCT, ""),
    ("sell_land", "Selling costs - lots / land / pads", 0.04, PCT, ""),
    ("sell_inc", "Sale costs - stabilized income property", 0.02, PCT, ""),
    ("taxrate", "Property tax rate", 0.022, PCT, "Texas, all jurisdictions combined - verify"),
    ("assess", "Assessed value as % of cost", 0.85, PCT, "Applied to total project cost"),
    ("ltc", "Development loan-to-cost", 0.65, PCT, ""),
    ("rate", "Development loan rate", 0.09, PCT, ""),
    ("pts", "Development loan points", 0.01, PCT, ""),
    ("bal", "Average balance drawn during construction", 0.55, PCT, ""),
    ("leaseint", "Share of lease-up interest not covered by ramping NOI", 0.50, PCT, "Income uses only"),
    ("pltv", "Permanent loan max LTV", 0.70, PCT, ""),
    ("pdscr", "Permanent loan min DSCR", 1.25, MULT, ""),
    ("pdy", "Permanent loan min debt yield", 0.09, PCT, ""),
    ("prate", "Permanent loan rate", 0.068, PCT, ""),
    ("pamort", "Permanent loan amortization", 30, NUM, "years"),
    ("raw", "Value of excess land not used by the use", 12000, USD, "$ per gross acre; resale of leftover acres"),
    ("disc", "Opening offer discount below max", 0.15, PCT, ""),
]
for i, (k, lab, v, fmt, note) in enumerate(alist):
    r = 5 + i; A[k] = f"Assumptions!$C${r}"
    put(asn, f"B{r}", lab); put(asn, f"C{r}", v, fin, fmt, fill_in); put(asn, f"D{r}", note, fs)
asn.column_dimensions["A"].width = 2; asn.column_dimensions["B"].width = 46; asn.column_dimensions["C"].width = 14; asn.column_dimensions["D"].width = 46
# parcel utility logic needs A
pc[P['pubW'].split('$')[1] + P['pubW'].split('$')[2]] = f'=IF(OR({P["water_site"]}="Y",{P["water_dist"]}<={A["wmax"]}),1,0)'
pc[P['pubS'].split('$')[1] + P['pubS'].split('$')[2]] = f'=IF(OR({P["sewer_site"]}="Y",{P["sewer_dist"]}<={A["smax"]}),1,0)'
pc[P['extW'].split('$')[1] + P['extW'].split('$')[2]] = f'=IF({P["water_site"]}="Y",0,{P["water_dist"]})*{P["pubW"]}'
pc[P['extS'].split('$')[1] + P['extS'].split('$')[2]] = f'=IF({P["sewer_site"]}="Y",0,{P["sewer_dist"]})*{P["pubS"]}'

# ============================================================ Cost library
cl = wb.create_sheet("Cost_Library", 2)
title(cl, "Cost Library - horizontal unit costs", "PLACEHOLDERS. In the software this is built from ingested bids, proformas and actuals (see build guide).")
hdr(cl, 4, ["Code", "Line item", "Unit", "Low", "Base", "High", "Used (scenario)", "Adjustment", "Source quality", "Data points"], 1)
lib = [
    ("02-100", "Boundary & topo survey", "LS", 4000, 6500, 10000, "Regional"),
    ("02-110", "Platting / lot surveys", "per lot", 600, 900, 1400, "Regional"),
    ("02-200", "Civil engineering", "per developed ac", 3000, 5000, 8000, "Regional"),
    ("02-300", "Geotechnical report", "LS", 3000, 5000, 9000, "Regional"),
    ("02-310", "Septic site evaluations", "per system", 350, 500, 800, "Regional"),
    ("02-400", "Phase I environmental", "LS", 2000, 2800, 4000, "None"),
    ("02-500", "Drainage study", "LS", 4000, 7500, 15000, "Regional"),
    ("02-700", "Legal, title & entity", "LS", 5000, 10000, 20000, "None"),
    ("03-100", "Clearing & grubbing", "per ac", 2500, 4000, 6500, "Regional"),
    ("03-200", "Mass grading", "per ac", 3000, 5000, 8000, "Regional x slope"),
    ("03-300", "Internal roads / drives", "per LF", 110, 150, 210, "Regional"),
    ("03-350", "Parking / paving / hardscape", "per SF", 4.5, 6, 8, "Regional"),
    ("03-400", "Detention pond", "per ac", 35000, 55000, 85000, "Regional"),
    ("03-500", "Erosion control / SWPPP", "per net ac", 400, 700, 1100, "Regional"),
    ("04-100", "Water line", "per LF", 35, 50, 75, "Regional"),
    ("04-110", "Water taps & meters", "per tap", 1500, 2500, 4000, "None"),
    ("04-200", "Sewer line", "per LF", 60, 85, 120, "Regional"),
    ("04-210", "Sewer taps / impact fees", "per tap", 1500, 3000, 5000, "None"),
    ("04-300", "On-site septic systems", "per system", 8000, 11000, 16000, "Regional x soil"),
    ("04-310", "Private water wells", "per well", 9000, 13000, 20000, "Regional"),
    ("04-400", "Electric extension", "per LF", 15, 25, 40, "Regional"),
]
C = {}
for i, (code, item, unit, lo, ba, hi, adj) in enumerate(lib):
    r = 5 + i; C[code] = f"Cost_Library!$G${r}"
    put(cl, f"A{r}", code); put(cl, f"B{r}", item); put(cl, f"C{r}", unit)
    fmt = USD2 if ba < 10 else USD
    for col, v in zip("DEF", (lo, ba, hi)): put(cl, f"{col}{r}", v, fin, fmt, fill_in)
    put(cl, f"G{r}", f"=INDEX(D{r}:F{r},{P['sidx']})", fn, fmt); put(cl, f"H{r}", adj)
    put(cl, f"I{r}", "Placeholder", fin); put(cl, f"J{r}", 0, fin, NUM)
for k, w in {"A": 9, "B": 32, "C": 18, "D": 10, "E": 10, "F": 10, "G": 14, "H": 16, "I": 14, "J": 11}.items():
    cl.column_dimensions[k].width = w
cl.freeze_panes = "C5"

# ============================================================ Use library
ul = wb.create_sheet("Use_Library", 3)
title(ul, "Use Library", "One row per use. Every economic value is a PLACEHOLDER for North Texas - replace with comps, bids and your own actuals. Add a use by adding a row (and extending the Engine tab).")
hdr(ul, 4, [lab for _, lab in COLS], 1, height=60)
FIRST = 5; N = len(U); LAST = FIRST + N - 1
UC = {k: L(i + 1) for i, k in enumerate(KEYS)}
pct_keys = {"os", "othinc", "vac", "mgmt", "recov", "cap", "margin"}
usd_keys = {"entitle", "vert", "lotprice", "saleprice", "rent", "opex", "tilc", "landac", "groundrent"}
for i, u in enumerate(U):
    r = FIRST + i
    for k in KEYS:
        v = u[k]; c = ul[f"{UC[k]}{r}"]; c.value = v
        if k in ("use", "cat"): c.font = fb if k == "use" else fn
        elif k in ("loc", "notes"): c.font = fs
        else: c.font = fin; c.fill = fill_in
        if k in pct_keys: c.number_format = PCT
        elif k in usd_keys: c.number_format = USD2 if isinstance(v, float) and v < 100 else USD
        elif k in ("density", "maxac", "minac"): c.number_format = '0.0#'
        elif k in ("sitesf", "road", "pave", "taps", "septic"): c.number_format = '#,##0.0###'
widths = {"corner": 10, "use": 40, "cat": 18, "engine": 13, "unit": 15, "loc": 60, "notes": 50}
for k in KEYS: ul.column_dimensions[UC[k]].width = widths.get(k, 12)
ul.freeze_panes = "B5"
dv(ul, f"{UC['engine']}{FIRST}:{UC['engine']}{LAST+40}", "=Lookups!$N$5:$N$8")
for k in ("discrete", "lotrule", "reqsewer", "reqwater", "horiz"):
    dv(ul, f"{UC[k]}{FIRST}:{UC[k]}{LAST+40}", "=Lookups!$L$5:$L$6")

# ============================================================ Engine (row per use)
en = wb.create_sheet("Engine", 4)
title(en, "Engine - every use run against this parcel", "One row per use, one column per calculation step, left to right. This tab IS the spec: each column header is a step Ryan codes. Do not edit - change inputs on Parcel, Assumptions, Use_Library.")
ENG = []  # (key, header, formula template, fmt, group)
def col(key, header, f, fmt=USD, grp=""):
    ENG.append((key, header, f, fmt, grp))

G = "Identity"
col("use", "Use", "={u:use}", None, G); col("engine", "Engine", "={u:engine}", None, G)
col("unit", "Capacity unit", "={u:unit}", None, G)
G = "Yield"
col("usable_all", "Usable acres (net - open space - detention)", "=MAX(0,{p:net}*(1-{u:os}-{a:det}))", NUM1, G)
col("used_ac", "Usable acres used by this use (market cap)", "=IF({u:maxac}>0,MIN({e:usable_all},{u:maxac}),{e:usable_all})", NUM1, G)
col("ufrac", "Share of parcel used", "=IF({e:usable_all}=0,0,{e:used_ac}/{e:usable_all})", PCT, G)
col("net_u", "Net acres in use", "={p:net}*{e:ufrac}", NUM1, G)
col("det_ac", "Detention acres", "={e:net_u}*{a:det}", NUM1, G)
col("dev_ac", "Developed acres", "={e:net_u}*(1-{u:os})", NUM1, G)
col("excess_ac", "Excess gross acres (resold as raw land)", "={p:gross}*(1-{e:ufrac})", NUM1, G)
col("sewer_ok", "Sewer requirement met (1/0)", '=IF(AND({u:reqsewer}="Y",{p:pubS}=0),0,1)', NUM, G)
col("water_ok", "Water requirement met (1/0)", '=IF(AND({u:reqwater}="Y",{p:pubW}=0),0,1)', NUM, G)
col("loc_fail", "Location test failures", '=TRIM(IF(AND({u:minvpd}>0,{p:vpd}<{u:minvpd}),"Traffic under "&TEXT({u:minvpd},"#,##0")&" VPD. ","")&IF(AND({u:minhh}>0,{p:hh}<{u:minhh}),"Under "&TEXT({u:minhh},"#,##0")&" households in 3 mi. ","")&IF(AND({u:maxhwy}>0,{p:hwy}>{u:maxhwy}),"Over "&{u:maxhwy}&" mi to highway. ","")&IF(AND({u:corner}="Y",{p:corner}<>"Y"),"Needs a corner. ","")&IF(AND({u:maxsub}>0,{p:sub}>{u:maxsub}),"Over "&{u:maxsub}&" mi to substation. ",""))', None, G)
col("loc_ok", "Location passes (1/0)", '=IF({e:loc_fail}="",1,0)', NUM, G)
col("size_ok", "Parcel big enough (1/0)", "=IF({e:usable_all}>={u:minac},1,0)", NUM, G)
col("cap_den", "Capacity by density", "=IF({u:density}>0,{e:net_u}*{u:density},1E+9)", NUM1, G)
col("cap_site", "Capacity by site area", "=IF({u:sitesf}>0,{e:used_ac}*43560/{u:sitesf},1E+9)", NUM1, G)
col("cap_lot", "Capacity by septic lot size", '=IF(AND({u:lotrule}="Y",{p:pubS}=0),{e:used_ac}*(1-{a:rowpct})/IF({p:pubW}=1,{a:lot_w},{a:lot_well}),1E+9)', NUM1, G)
col("capacity", "CAPACITY (units / SF / acres)", '=IF({e:sewer_ok}*{e:water_ok}*{e:size_ok}*{e:loc_ok}=0,0,IF(MIN({e:cap_den},{e:cap_site},{e:cap_lot})>=1E+9,{e:used_ac},IF({u:discrete}="Y",INT(MIN({e:cap_den},{e:cap_site},{e:cap_lot})),MIN({e:cap_den},{e:cap_site},{e:cap_lot}))))', NUM, G)
col("binding", "Binding constraint", '=IF({e:loc_ok}=0,"Location fails",IF({e:size_ok}=0,"Too small - needs "&{u:minac}&" usable ac",IF({e:sewer_ok}=0,"Needs public sewer",IF({e:water_ok}=0,"Needs public water",IF({e:cap_lot}<=MIN({e:cap_den},{e:cap_site}),"Septic lot size",IF({e:cap_site}<={e:cap_den},"Site area / parking","Density"))))))', None, G)
col("on", "Anything fits (1/0)", "=IF({e:capacity}>0,1,0)", NUM, G)
col("hz", "Horizontal by us (1/0)", '=IF({u:horiz}="Y",1,0)*{e:on}', NUM, G)
G = "Quantities"
col("road_lf", "Internal road LF", "={e:capacity}*{u:road}*{e:hz}", NUM, G)
col("pave_sf", "Paving / hardscape SF", "={e:capacity}*{u:pave}*{e:hz}", NUM, G)
col("water_lf", "Water line LF", "=IF({p:pubW}=1,{e:road_lf}+{p:extW},0)*{e:hz}", NUM, G)
col("wtaps", "Water taps", "={e:capacity}*{u:taps}*{p:pubW}*{e:hz}", NUM1, G)
col("sewer_lf", "Sewer line LF", "=IF({p:pubS}=1,{e:road_lf}+{p:extS},0)*{e:hz}", NUM, G)
col("staps", "Sewer taps", "={e:capacity}*{u:taps}*{p:pubS}*{e:hz}", NUM1, G)
col("septics", "Septic systems", "=IF({p:pubS}=0,{e:capacity}*{u:septic},0)*{e:hz}", NUM1, G)
col("wells", "Private wells", "=IF({p:pubW}=0,{e:capacity}*{u:septic},0)*{e:hz}", NUM1, G)
col("clear_ac", "Acres to clear", "={e:dev_ac}*{p:tree}*{e:hz}", NUM1, G)
col("grade_ac", "Acres to grade", "={e:dev_ac}*{e:hz}", NUM1, G)
col("plats", "Lots to plat", '=IF({u:engine}="For-Sale",IF({u:discrete}="Y",{e:capacity},1),1)*{e:on}', NUM, G)
G = "Horizontal budget"
col("c_survey", "02-100 Survey", "={c:02-100}*{p:reg}*{e:on}", USD, G)
col("c_plat", "02-110 Platting", "={e:plats}*{c:02-110}*{p:reg}", USD, G)
col("c_civil", "02-200 Civil engineering", "={e:dev_ac}*{c:02-200}*{p:reg}*{e:hz}", USD, G)
col("c_geo", "02-300 Geotech", "={c:02-300}*{p:reg}*{e:hz}", USD, G)
col("c_septev", "02-310 Septic evaluations", "={e:septics}*{c:02-310}*{p:reg}", USD, G)
col("c_p1", "02-400 Phase I", "={c:02-400}*{e:on}", USD, G)
col("c_drain", "02-500 Drainage study", "={c:02-500}*{p:reg}*{e:hz}", USD, G)
col("c_ent", "Entitlement & zoning (use-specific)", "={u:entitle}*{e:on}", USD, G)
col("c_legal", "02-700 Legal & title", "={c:02-700}*{e:on}", USD, G)
col("c_clear", "03-100 Clearing", "={e:clear_ac}*{c:03-100}*{p:reg}", USD, G)
col("c_grade", "03-200 Grading", "={e:grade_ac}*{c:03-200}*{p:reg}*{p:slopem}", USD, G)
col("c_road", "03-300 Roads / drives", "={e:road_lf}*{c:03-300}*{p:reg}", USD, G)
col("c_pave", "03-350 Paving / parking", "={e:pave_sf}*{c:03-350}*{p:reg}", USD, G)
col("c_det", "03-400 Detention", "={e:det_ac}*{c:03-400}*{p:reg}*{e:hz}", USD, G)
col("c_swppp", "03-500 SWPPP", "={e:net_u}*{c:03-500}*{p:reg}*{e:hz}", USD, G)
col("c_wline", "04-100 Water line", "={e:water_lf}*{c:04-100}*{p:reg}", USD, G)
col("c_wtap", "04-110 Water taps", "={e:wtaps}*{c:04-110}", USD, G)
col("c_sline", "04-200 Sewer line", "={e:sewer_lf}*{c:04-200}*{p:reg}", USD, G)
col("c_stap", "04-210 Sewer taps", "={e:staps}*{c:04-210}", USD, G)
col("c_septic", "04-300 Septic systems", "={e:septics}*{c:04-300}*{p:reg}*{p:soilm}", USD, G)
col("c_well", "04-310 Wells", "={e:wells}*{c:04-310}*{p:reg}", USD, G)
col("c_elec", "04-400 Electric", "={e:road_lf}*{c:04-400}*{p:reg}", USD, G)
col("h_sub", "Horizontal subtotal", "=SUM({e:c_survey}:{e:c_elec})", USD, G)
col("h_cont", "Horizontal contingency", "={e:h_sub}*{a:hcont}", USD, G)
col("H", "HORIZONTAL TOTAL (ex land)", "={e:h_sub}+{e:h_cont}", USD, G)
G = "Vertical"
col("vert", "Vertical hard cost", "={e:capacity}*{u:vert}", USD, G)
col("vsc", "Vertical soft + contingency", "={e:vert}*({a:vsoft}+{a:vcont})", USD, G)
col("tilc", "TI + leasing commissions", "={e:capacity}*{u:tilc}", USD, G)
G = "Land & financing factors"
col("landfac", "Land cost factor (closing + carry)", "=1+{a:closing}+{a:carry}*{u:devmo}/12", NUM3, G)
col("ff_dev", "Financing factor - no lease-up", "=1+{a:ltc}*({a:rate}*{u:devmo}/12*{a:bal}+{a:pts})", NUM3, G)
col("ff_lease", "Financing factor - with lease-up", '=1+{a:ltc}*({a:rate}*({u:devmo}/12*{a:bal}+IF({u:engine}="Income",{u:leaseup},0)/12*{a:leaseint})+{a:pts})', NUM3, G)
col("excess_rev", "Excess land resale (net)", "={e:excess_ac}*{a:raw}*(1-{a:sell_land})*{e:on}", USD, G)
G = "Exit 1 - Land / pad sale"
col("rev1", "Exit 1 net revenue", '=IF({u:ex1}=1,IF({u:engine}="For-Sale",{e:capacity}*{u:lotprice},{e:used_ac}*{u:landac})*(1-{a:sell_land})+{e:excess_rev},0)*{e:on}', USD, G)
col("nl1", "Exit 1 non-land cost (no parking paving: buyer builds it)", "=({e:h_sub}-{e:c_pave})*(1+{a:hcont})", USD, G)
col("k1", "Exit 1 cost per $1 of land", "={e:landfac}*{e:ff_dev}", NUM3, G)
col("cost1", "Exit 1 total cost at asking", "=({p:land}*{e:landfac}+{e:nl1})*{e:ff_dev}", USD, G)
col("profit1", "Exit 1 profit at asking", "={e:rev1}-{e:cost1}", USD, G)
col("max1", "EXIT 1 MAX OFFER", "=IF(AND({u:ex1}=1,{e:on}=1),MAX(0,((1-{u:margin})*{e:rev1}-{e:nl1}*{e:ff_dev})/{e:k1}),0)", USD, G)
G = "Build cost (exits 2 & 3)"
col("nl2", "Build non-land cost (pre-financing)", "={e:H}+{e:vert}+{e:vsc}+{e:tilc}", USD, G)
col("k2", "Build cost per $1 of land", "={e:landfac}*{e:ff_lease}", NUM3, G)
col("cost2", "Build total cost at asking", "=({p:land}*{e:landfac}+{e:nl2})*{e:ff_lease}", USD, G)
G = "Income (Income engine)"
inc = '{u:engine}="Income"'
col("gpr", "Gross potential rent", f"=IF({inc},{{e:capacity}}*{{u:rent}},0)", USD, G)
col("egi", "Effective gross income", "={e:gpr}*(1+{u:othinc})*(1-{u:vac})", USD, G)
col("opex", "Operating expenses (ex tax, ex mgmt)", f"=IF({inc},{{e:capacity}}*{{u:opex}},0)", USD, G)
col("ptax", "Property tax", f"=IF({inc},{{a:taxrate}}*{{a:assess}}*{{e:cost2}},0)", USD, G)
col("mgmtx", "Management", "={e:egi}*{u:mgmt}", USD, G)
col("recov", "Expense recoveries", "=({e:opex}+{e:ptax})*{u:recov}", USD, G)
col("noi", "NET OPERATING INCOME", "={e:egi}-{e:opex}-{e:ptax}-{e:mgmtx}+{e:recov}", USD, G)
col("value", "Stabilized value (NOI / cap)", f"=IF(AND({inc},{{u:cap}}>0),{{e:noi}}/{{u:cap}},0)", USD, G)
col("dnoi", "Change in NOI per $1 of land", f"=IF({inc},-{{a:taxrate}}*{{a:assess}}*(1-{{u:recov}})*{{e:k2}},0)", '0.0000', G)
col("yoc", "Yield on cost", "=IF({e:cost2}=0,0,{e:noi}/{e:cost2})", PCT, G)
G = "Exit 2 - Build & sell"
col("rev2", "Exit 2 net revenue", f'=IF({{u:ex2}}=1,IF({inc},{{e:value}}*(1-{{a:sell_inc}}),{{e:capacity}}*{{u:saleprice}}*(1-{{a:sell_fs}}))+{{e:excess_rev}},0)*{{e:on}}', USD, G)
col("drev2", "Change in exit 2 revenue per $1 of land", f'=IF(AND({{u:ex2}}=1,{inc},{{u:cap}}>0),{{e:dnoi}}/{{u:cap}}*(1-{{a:sell_inc}}),0)', '0.0000', G)
col("profit2", "Exit 2 profit at asking", "={e:rev2}-{e:cost2}", USD, G)
col("max2", "EXIT 2 MAX OFFER", "=IF(AND({u:ex2}=1,{e:on}=1),MAX(0,{p:land}-((1-{u:margin})*{e:rev2}-{e:cost2})/((1-{u:margin})*{e:drev2}-{e:k2})),0)", USD, G)
G = "Exit 3 - Build & hold / lease"
col("rev3", "Exit 3 value", f'=IF({{u:ex3}}=1,IF({inc},{{e:value}},IF({{u:cap}}>0,{{e:used_ac}}*{{u:groundrent}}/{{u:cap}},0))+{{e:excess_rev}},0)*{{e:on}}', USD, G)
col("nl3", "Exit 3 non-land cost", f"=IF({inc},{{e:nl2}},{{e:H}})", USD, G)
col("k3", "Exit 3 cost per $1 of land", f"=IF({inc},{{e:k2}},{{e:landfac}}*{{e:ff_dev}})", NUM3, G)
col("cost3", "Exit 3 total cost at asking", f"=IF({inc},{{e:cost2}},({{p:land}}*{{e:landfac}}+{{e:H}})*{{e:ff_dev}})", USD, G)
col("drev3", "Change in exit 3 value per $1 of land", f'=IF(AND({{u:ex3}}=1,{inc},{{u:cap}}>0),{{e:dnoi}}/{{u:cap}},0)', '0.0000', G)
col("profit3", "Exit 3 value created at asking", "={e:rev3}-{e:cost3}", USD, G)
col("max3", "EXIT 3 MAX OFFER", "=IF(AND({u:ex3}=1,{e:on}=1),MAX(0,{p:land}-((1-{u:margin})*{e:rev3}-{e:cost3})/((1-{u:margin})*{e:drev3}-{e:k3})),0)", USD, G)
G = "Hold financing (Income)"
col("devloan", "Development loan", "=IF({inc_q},{a:ltc}*({p:land}*{e:landfac}+{e:nl2}),0)".replace("{inc_q}", inc), USD, G)
col("ploan", "Permanent loan (lesser of LTV, DSCR, debt yield)", f"=IF(AND({inc},{{e:noi}}>0),MIN({{a:pltv}}*{{e:value}},PV({{a:prate}}/12,{{a:pamort}}*12,-{{e:noi}}/{{a:pdscr}}/12),{{e:noi}}/{{a:pdy}}),0)", USD, G)
col("dscr", "DSCR on permanent loan", "=IF({e:ploan}=0,0,{e:noi}/(-PMT({a:prate}/12,{a:pamort}*12,{e:ploan})*12))", MULT, G)
col("cashout", "Cash out at refi (perm - dev loan)", "={e:ploan}-{e:devloan}", USD, G)
G = "Answer"
col("best", "BEST MAX OFFER", "=MAX({e:max1},{e:max2},{e:max3})", USD, G)
col("bestexit", "Best exit", '=IF({e:best}<=0,"None works",IF({e:max1}={e:best},"Land / pad sale",IF({e:max2}={e:best},"Build & sell","Build & hold / lease")))', None, G)
col("open", "Opening offer", "={e:best}*(1-{a:disc})", USD, G)
col("spread", "Max offer minus asking", "={e:best}-{p:land}", USD, G)
col("peracre", "Max offer per gross acre", "=IF({p:gross}=0,0,{e:best}/{p:gross})", USD, G)
col("perunit", "Max offer per unit of capacity", "=IF({e:capacity}=0,0,{e:best}/{e:capacity})", USD2, G)
col("flags", "Flags", '=TRIM({e:loc_fail}&" "&IF(AND({e:on}=0,{e:loc_ok}=1),"Nothing fits. ","")&IF({p:frontage}<{u:minfront},"Frontage under "&{u:minfront}&" LF. ","")&IF({p:slope}>0.15,"Slope over 15%. ","")&IF(AND({e:on}=1,{u:maxac}>0,{e:usable_all}>{u:maxac}),"Only "&{u:maxac}&" ac used - rest resold. ",""))', None, G)
col("rankkey", "Rank key", "={e:best}+ROW()/1000000", USD, G)

EC = {k: L(i + 1) for i, (k, *_ ) in enumerate(ENG)}
def expand(f, ur, r):
    f = re.sub(r"\{u:(\w+)\}", lambda m: f"Use_Library!${UC[m.group(1)]}${ur}", f)
    f = re.sub(r"\{e:(\w+)\}", lambda m: f"{EC[m.group(1)]}{r}", f)
    f = re.sub(r"\{p:(\w+)\}", lambda m: P[m.group(1)], f)
    f = re.sub(r"\{a:(\w+)\}", lambda m: A[m.group(1)], f)
    f = re.sub(r"\{c:([\d-]+)\}", lambda m: C[m.group(1)], f)
    assert "{" not in f, f
    return f

# group header row 4, column header row 5
prev = None
for i, (k, h, f, fmt, grp) in enumerate(ENG):
    c = i + 1
    if grp != prev:
        x = en.cell(row=4, column=c, value=grp); x.font = Font(name=F, size=10, bold=True, color=NAVY); prev = grp
hdr(en, 5, [h for _, h, *_ in ENG], 1, height=70)
EFIRST = 6
for j in range(N):
    r = EFIRST + j; ur = FIRST + j
    for i, (k, h, f, fmt, grp) in enumerate(ENG):
        cell = en.cell(row=r, column=i + 1, value=expand(f, ur, r))
        cell.font = fb if k in ("capacity", "best", "use") else fn
        if fmt: cell.number_format = fmt
        if k in ("max1", "max2", "max3", "best"): cell.fill = fill_g
ELAST = EFIRST + N - 1
for i, (k, *_ ) in enumerate(ENG):
    en.column_dimensions[L(i + 1)].width = 38 if k == "use" else (30 if k in ("flags",) else 14)
en.freeze_panes = "B6"
ER = {k: f"Engine!${EC[k]}${EFIRST}:${EC[k]}${ELAST}" for k in EC}

# ============================================================ Best use ranking
bu = wb.create_sheet("Best_Use", 0)
title(bu, "Best Use - what this parcel is worth to each use", "Every use in the library run against the parcel on the Parcel tab, ranked by max offer. This is the 'land value today to its best use' screen.")
put(bu, "B3", "Parcel:", fb); put(bu, "C3", f"={P['name']}", flink)
put(bu, "E3", "Asking:", fb); put(bu, "F3", f"={P['land']}", flink, USD)
hdr(bu, 5, ["Rank", "Use", "Engine", "Capacity", "Unit", "Best exit", "Max offer", "Opening offer", "Max - asking", "Max / gross ac", "Max / unit", "Binding constraint", "Flags"], 1, height=32)
for j in range(N):
    r = 6 + j
    put(bu, f"A{r}", j + 1, fn, NUM)
    put(bu, f"N{r}", f"=MATCH(LARGE({ER['rankkey']},A{r}),{ER['rankkey']},0)", fn, NUM)
    for colL, k, fmt in [("B", "use", None), ("C", "engine", None), ("D", "capacity", NUM), ("E", "unit", None), ("F", "bestexit", None),
                         ("G", "best", USD), ("H", "open", USD), ("I", "spread", USD), ("J", "peracre", USD), ("K", "perunit", USD),
                         ("L", "binding", None), ("M", "flags", None)]:
        c = put(bu, f"{colL}{r}", f"=INDEX({ER[k]},$N{r})", fb if k in ("use", "best") else fn, fmt)
        if k == "best": c.fill = fill_g
put(bu, "N5", "row", fs)
for k, w in {"A": 6, "B": 42, "C": 13, "D": 11, "E": 14, "F": 20, "G": 14, "H": 14, "I": 14, "J": 13, "K": 12, "L": 22, "M": 52, "N": 5}.items():
    bu.column_dimensions[k].width = w
bu.freeze_panes = "C6"
put(bu, f"B{7+N}", "Read with care: every number is a placeholder, and a high max offer means only that the math works. Check the location requirements (Use_Library) before believing any row.", fs)

# ============================================================ Deal view (one use, top to bottom)
dvw = wb.create_sheet("Deal_View", 1)
title(dvw, "Deal View - one use in detail", "Pick a use. Everything below is read from the Engine row for that use, top to bottom, so you can audit each step.")
put(dvw, "B4", "Use", fb); put(dvw, "C4", "Multifamily - large garden (150+ units)", fin, None, fill_in)
dv(dvw, "C4", f"=Use_Library!$A${FIRST}:$A${LAST}")
put(dvw, "B5", "Engine row", fs); put(dvw, "C5", f"=MATCH(C4,{ER['use']},0)", fn, NUM)
put(dvw, "B6", "Location requirements", fb)
put(dvw, "C6", f"=INDEX(Use_Library!${UC['loc']}${FIRST}:${UC['loc']}${LAST},$C$5)", fs)
r = 8; prev = None
skip = {"use", "rankkey", "max1", "max2", "max3", "rev1", "rev2", "rev3", "cost1", "cost2", "cost3", "profit1", "profit2", "profit3",
        "nl1", "nl3", "k1", "k3", "drev2", "drev3", "best", "bestexit", "open", "spread", "peracre", "perunit", "flags"}
for k, h, f, fmt, grp in ENG:
    if k in skip: continue
    if grp != prev:
        r += 1; put(dvw, f"B{r}", grp, fsec); r += 1; prev = grp
    put(dvw, f"B{r}", h); c = put(dvw, f"C{r}", f"=INDEX({ER[k]},$C$5)", fb if k in ("capacity", "H", "noi") else fn, fmt)
    r += 1
r += 1
put(dvw, f"B{r}", "Exits", fsec); r += 1
hdr(dvw, r, ["Line", "Land / pad sale", "Build & sell", "Build & hold / lease"], 2); r += 1
for lab, ks, fmt in [("Allowed for this use", ("ex1", "ex2", "ex3"), None),
                     ("Net revenue / value", ("rev1", "rev2", "rev3"), USD),
                     ("Total cost at asking price", ("cost1", "cost2", "cost3"), USD),
                     ("Profit / value created at asking", ("profit1", "profit2", "profit3"), USD),
                     ("Cost per $1 of land", ("k1", "k2", "k3"), NUM3),
                     ("MAX OFFER", ("max1", "max2", "max3"), USD)]:
    put(dvw, f"B{r}", lab, fb if lab == "MAX OFFER" else fn)
    for colL, k in zip("CDE", ks):
        if k.startswith("ex"):
            put(dvw, f"{colL}{r}", f'=IF(INDEX(Use_Library!${UC[k]}${FIRST}:${UC[k]}${LAST},$C$5)=1,"Yes","No")')
        else:
            c = put(dvw, f"{colL}{r}", f"=INDEX({ER[k]},$C$5)", fb if lab == "MAX OFFER" else fn, fmt)
            if lab == "MAX OFFER": c.fill = fill_g
    r += 1
put(dvw, f"B{r}", "Margin at asking price", fn)
for colL, a, b in zip("CDE", ("profit1", "profit2", "profit3"), ("rev1", "rev2", "rev3")):
    put(dvw, f"{colL}{r}", f"=IF(INDEX({ER[b]},$C$5)=0,0,INDEX({ER[a]},$C$5)/INDEX({ER[b]},$C$5))", fn, PCT)
r += 1
put(dvw, f"B{r}", "Required margin", fn); put(dvw, f"C{r}", f"=INDEX(Use_Library!${UC['margin']}${FIRST}:${UC['margin']}${LAST},$C$5)", flink, PCT); r += 2
for lab, k, fmt in [("BEST MAX OFFER (walk-away)", "best", USD), ("Best exit", "bestexit", None), ("Opening offer", "open", USD),
                    ("Max offer minus asking", "spread", USD), ("Max offer per gross acre", "peracre", USD),
                    ("Max offer per unit", "perunit", USD), ("Flags", "flags", None)]:
    put(dvw, f"B{r}", lab, fb); c = put(dvw, f"C{r}", f"=INDEX({ER[k]},$C$5)", fb, fmt); c.fill = fill_g; r += 1
dvw.column_dimensions["A"].width = 2; dvw.column_dimensions["B"].width = 48
for k in "CDE": dvw.column_dimensions[k].width = 22

# ============================================================ Deal history + import mapping
dh = wb.create_sheet("Deal_History")
title(dh, "Deal History - upload template", "Past deals with projected AND actual numbers, including deals you passed on. This is what lets the system learn. Row 6 is an example.")
hcols = [("Deal_ID", "Text"), ("Date_Underwritten", "YYYY-MM-DD"), ("County", "Text"), ("Use", "From Use_Library"),
         ("Gross_Acres", "ac"), ("Capacity", "#"), ("Land_Price", "$"), ("Decision", "Bought / Passed / Lost"),
         ("Reason_if_Passed", "Text"), ("Horizontal_Projected", "$"), ("Horizontal_Actual", "$"), ("Vertical_Projected", "$"),
         ("Vertical_Actual", "$"), ("Revenue_or_Value_Projected", "$"), ("Revenue_or_Value_Actual", "$"),
         ("NOI_Projected", "$"), ("NOI_Actual", "$"), ("Months_Projected", "#"), ("Months_Actual", "#"),
         ("Data_Quality", "Actual / Bid / Estimate"), ("Notes", "Text"), ("Cost_Variance_%", "calc"), ("Revenue_Variance_%", "calc")]
hdr(dh, 4, [c for c, _ in hcols], 1, height=32)
for i, (_, t) in enumerate(hcols): put(dh, f"{L(i+1)}5", t, fs)
ex = ["EX-001", "2025-03-14", "Grayson", "Manufactured home subdivision", 12, 7, 140000, "Bought", "", 177000, 188500, 815500,
      830000, 1715000, 1690000, None, None, 14, 19, "Actual", "Example row - delete before upload"]
for i, v in enumerate(ex):
    c = dh.cell(row=6, column=i + 1, value=v); c.font = fin
    if isinstance(v, int) and v > 1000: c.number_format = USD
HI = {c: L(i + 1) for i, (c, _) in enumerate(hcols)}
for r in range(6, 106):
    put(dh, f"{HI['Cost_Variance_%']}{r}", f'=IF(OR({HI["Horizontal_Actual"]}{r}="",{HI["Horizontal_Projected"]}{r}=""),"",({HI["Horizontal_Actual"]}{r}+N({HI["Vertical_Actual"]}{r}))/({HI["Horizontal_Projected"]}{r}+N({HI["Vertical_Projected"]}{r}))-1)', fn, PCT)
    put(dh, f"{HI['Revenue_Variance_%']}{r}", f'=IF(OR({HI["Revenue_or_Value_Actual"]}{r}="",N({HI["Revenue_or_Value_Projected"]}{r})=0),"",{HI["Revenue_or_Value_Actual"]}{r}/{HI["Revenue_or_Value_Projected"]}{r}-1)', fn, PCT)
for i in range(len(hcols)): dh.column_dimensions[L(i + 1)].width = 16
dh.column_dimensions["D"].width = 30
dv(dh, f"D6:D105", f"=Use_Library!$A${FIRST}:$A${LAST}")
dv(dh, f"{HI['Decision']}6:{HI['Decision']}105", '"Bought,Passed,Lost to other buyer"')
dv(dh, f"{HI['Data_Quality']}6:{HI['Data_Quality']}105", '"Actual,Bid,Estimate"')
dh.freeze_panes = "B6"

im = wb.create_sheet("Import_Mapping")
title(im, "Import Mapping - how an uploaded spreadsheet lands in the schema", "Illustrative. The software proposes each mapping with a confidence score; the user confirms; confirmed mappings are remembered per company.")
hdr(im, 4, ["Uploaded file / tab", "Raw label", "Raw value", "Maps to", "Conversion", "Confidence", "Status", "Reason"], 1)
for i, row in enumerate([
    ("Sadler.xlsx / Budget", "Dirt work", 62000, "03-200 Mass grading", "Lump sum -> per acre", 0.86, "Confirmed", "Common synonym"),
    ("Sadler.xlsx / Budget", "Septics x7", 77000, "04-300 Septic systems", "/ 7 -> per system", 0.95, "Auto-accepted", "Quantity in label"),
    ("Sadler.xlsx / Budget", "Misc / other", 15000, "Unmapped", "", 0.22, "Needs review", "Too vague"),
    ("Flex_deal.xlsx / Proforma", "Base rent PSF", 14.5, "Use_Library rent (per bldg SF/yr)", "", 0.93, "Auto-accepted", "Unit matches"),
    ("Flex_deal.xlsx / Proforma", "CAM reimbursements", 98000, "Expense recovery %", "/ (opex + tax)", 0.71, "Needs review", "Derived ratio"),
    ("MHC_buy.xlsx / Rent roll", "Lot rent", 525, "Use_Library rent (per pad/yr)", "x 12", 0.90, "Auto-accepted", "Monthly -> annual"),
]):
    for j, v in enumerate(row):
        c = im.cell(row=5 + i, column=j + 1, value=v); c.font = fn
        if j == 5: c.number_format = PCT
for k, w in {"A": 28, "B": 22, "C": 12, "D": 32, "E": 22, "F": 11, "G": 15, "H": 20}.items(): im.column_dimensions[k].width = w

# ============================================================ ReadMe
rm = wb.create_sheet("ReadMe", 0)
title(rm, "Landova Proforma v2", "Every Landova use, one engine. Parcel facts + use defaults + cost library -> what fits -> cost -> three exits -> max offer.")
lines = [
    ("Start here", None),
    ("1. Parcel", "Enter or paste the parcel's facts and the asking price."),
    ("2. Best_Use", "Every use ranked by the most you could pay for this land. This is the 'land value today to its best use' screen."),
    ("3. Deal_View", "Pick one use and read its numbers top to bottom: yield, quantities, budget, income, exits, max offer."),
    ("", None),
    ("How it is built", None),
    ("Four engines", "For-Sale (lots, homes, condos) · Income (anything 5+ units or commercial: valued on NOI and cap rate) · Operator Pad (Landova delivers a pad, an operator runs the business) · Land Play (leases and entitle-and-sell)."),
    ("Three exits for every use", "1) Land / pad sale  2) Build & sell  3) Build & hold / lease. Use_Library says which are allowed."),
    ("Use_Library", "One row per use: how it is sized, what it costs, what it earns, and where it works. Add a use = add a row."),
    ("Engine", "One row per use, one column per step. This tab is the spec Ryan codes from - every column header is a calculation step."),
    ("Max offer", "The engine solves for the land price that leaves exactly the required margin. Costs that rise with land price (closing, carry, financing, property tax) are included, so the answer is exact."),
    ("Market cap", "Max usable acres stops the model from pretending 20 acres can all be QSR pads. Leftover acres are resold at the raw land value."),
    ("", None),
    ("Color code", None),
    ("Blue on yellow", "Input - edit"),
    ("Black", "Formula - don't edit"),
    ("Green", "Link to another tab"),
    ("", None),
    ("Before using on a live deal", None),
    ("Placeholders", "Every cost, price, rent, cap rate, county factor and land value is a placeholder chosen to show the mechanics. Replace with real comps, bids and your actuals."),
    ("Location tests", "Traffic, households, highway distance, corner and substation minimums in Use_Library are enforced: a use that fails gets no capacity and drops to the bottom. The written location notes are NOT enforced - the land search must apply them."),
    ("Not in this workbook", "Monthly cash flow / IRR, multifamily unit mix and LP/GP waterfall, operating-business P&Ls. See the build guide."),
]
r = 4
for a, b in lines:
    if b is None and a: put(rm, f"A{r}", a, fsec)
    elif a or b:
        put(rm, f"A{r}", a, fb); c = put(rm, f"B{r}", b); c.alignment = Alignment(wrap_text=True, vertical="top")
        if a == "Blue on yellow": rm[f"A{r}"].font = fin; rm[f"A{r}"].fill = fill_in
        if a == "Green": rm[f"A{r}"].font = flink
    r += 1
rm.column_dimensions["A"].width = 26; rm.column_dimensions["B"].width = 120

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
order = ["ReadMe", "Parcel", "Best_Use", "Deal_View", "Use_Library", "Engine", "Assumptions", "Cost_Library", "Deal_History", "Import_Mapping", "Lookups"]
wb._sheets = [wb[n] for n in order]
wb.save("/home/claude/landova/Landova_Proforma_v2.xlsx")
print("uses", N, "engine cols", len(ENG), "rows", EFIRST, ELAST)
import json; json.dump({"EC": EC, "EFIRST": EFIRST, "N": N}, open("/home/claude/landova/meta.json", "w"))
